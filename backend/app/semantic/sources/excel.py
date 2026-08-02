"""XLSX field description adapter."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.contracts.semantic import FieldDescriptionRow
from app.semantic.sanitizer import DescriptionSanitizer
from app.semantic.sources.base import (
    KnowledgeSourceError,
    build_description_row,
    resolve_columns,
)


class ExcelFieldDescriptionAdapter:
    """Read the minimum field_path + description contract from XLSX."""

    def __init__(
        self,
        source_path: Path,
        *,
        sheet_name: str | None = None,
        sanitizer: DescriptionSanitizer | None = None,
    ) -> None:
        self.source_path = source_path
        self.sheet_name = sheet_name
        self.sanitizer = sanitizer or DescriptionSanitizer()

    def load_field_descriptions(self) -> list[FieldDescriptionRow]:
        workbook = load_workbook(self.source_path, read_only=False, data_only=True)
        try:
            worksheet = self._select_worksheet(workbook.sheetnames, workbook)
            for merged_range in worksheet.merged_cells.ranges:
                if merged_range.max_row > 1:
                    raise KnowledgeSourceError(
                        f"merged business data cells are not allowed: {merged_range}"
                    )

            row_iterator = worksheet.iter_rows(values_only=True)
            header = next(row_iterator, None)
            if header is None:
                raise KnowledgeSourceError("field description workbook is empty")
            columns = resolve_columns(header)
            rows: list[FieldDescriptionRow] = []
            for source_row, values in enumerate(row_iterator, start=2):
                parsed = build_description_row(
                    values, columns, source_row=source_row, sanitizer=self.sanitizer
                )
                if parsed is not None:
                    rows.append(parsed)
            return rows
        finally:
            workbook.close()

    def _select_worksheet(self, sheet_names: list[str], workbook: Workbook) -> Worksheet:
        if self.sheet_name is not None:
            if self.sheet_name not in sheet_names:
                raise KnowledgeSourceError(f"worksheet does not exist: {self.sheet_name}")
            return workbook[self.sheet_name]
        for name in sheet_names:
            worksheet = workbook[name]
            if worksheet.max_row > 0 and worksheet.max_column > 0:
                return worksheet
        raise KnowledgeSourceError("workbook has no non-empty worksheet")
